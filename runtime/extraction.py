"""Bounded file inspection and text extraction for untrusted uploads."""
import io,zipfile

MAX_EXPANDED=50_000_000
MAX_ZIP_ENTRIES=2_000
MAX_IMAGE_PIXELS=40_000_000
TEXT_EXTENSIONS={'txt','md','csv','json','py','js','ts','html','xml','log'}
IMAGE_MIMES={'jpg':'image/jpeg','jpeg':'image/jpeg','png':'image/png','webp':'image/webp','gif':'image/gif'}
OFFICE_MIMES={
 'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
 'xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
 'pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation',
}

def _extension(name):
    value=str(name or '').strip().lower()
    if not value or '/' in value or '\\' in value or '.' not in value:raise ValueError('invalid filename')
    return value.rsplit('.',1)[-1]

def _inspect_zip(raw,ext):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        items=z.infolist()
        if len(items)>MAX_ZIP_ENTRIES:raise ValueError('too many archive entries')
        expanded=sum(i.file_size for i in items);compressed=sum(max(1,i.compress_size) for i in items)
        if expanded>MAX_EXPANDED or (expanded>5_000_000 and expanded/compressed>200):raise ValueError('expanded document too large')
        if any(i.flag_bits & 1 for i in items):raise ValueError('encrypted documents are not supported')
        names={i.filename for i in items};required={'docx':'word/document.xml','xlsx':'xl/workbook.xml','pptx':'ppt/presentation.xml'}[ext]
        if '[Content_Types].xml' not in names or required not in names:raise ValueError('file content does not match extension')

def inspect(raw,name,mime='application/octet-stream'):
    if not isinstance(raw,(bytes,bytearray)) or not raw:raise ValueError('empty file')
    ext=_extension(name);raw=bytes(raw)
    if ext in TEXT_EXTENSIONS:
        if b'\x00' in raw[:8192]:raise ValueError('binary data does not match text extension')
        return raw.decode('utf-8',errors='replace')[:1_000_000],('text/csv' if ext=='csv' else 'text/plain')
    if ext=='pdf':
        if not raw.lstrip().startswith(b'%PDF-'):raise ValueError('file content does not match PDF extension')
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(raw),strict=True)
        if len(reader.pages)>1_000:raise ValueError('PDF has too many pages')
        return '\n'.join((p.extract_text() or '') for p in reader.pages[:100])[:1_000_000],'application/pdf'
    if ext in OFFICE_MIMES:
        _inspect_zip(raw,ext)
        if ext=='docx':
            from docx import Document
            d=Document(io.BytesIO(raw));text='\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])
        elif ext=='xlsx':
            import openpyxl
            w=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True);parts=[]
            for s in w:
                parts.append(s.title)
                for row in s.iter_rows(max_row=2000,values_only=True):parts.append(' | '.join('' if x is None else str(x) for x in row))
            w.close();text='\n'.join(parts)
        else:
            from pptx import Presentation
            presentation=Presentation(io.BytesIO(raw));parts=[]
            for number,slide in enumerate(presentation.slides,1):
                parts.append('Slayt '+str(number))
                for shape in slide.shapes:
                    if hasattr(shape,'text') and shape.text.strip():parts.append(shape.text.strip())
            text='\n'.join(parts)
        return text[:1_000_000],OFFICE_MIMES[ext]
    if ext in IMAGE_MIMES:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=MAX_IMAGE_PIXELS
        try:
            image=Image.open(io.BytesIO(raw));image.verify()
        except Exception as exc:raise ValueError('invalid image') from exc
        expected={'jpg':'JPEG','jpeg':'JPEG','png':'PNG','webp':'WEBP','gif':'GIF'}[ext]
        if image.format!=expected:raise ValueError('file content does not match image extension')
        return '[Görsel dosyası: '+str(name)+']',IMAGE_MIMES[ext]
    raise ValueError('unsupported file type')

def extract(raw,name,mime):
    return inspect(raw,name,mime)[0]
