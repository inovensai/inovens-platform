"""Create polished text documents in the current user's isolated workspace."""
import html,re,urllib.parse
from datetime import date
from pathlib import Path
from common import ROOT,Denied

FORMATS={'pdf','.pdf','docx','.docx','md','.md','txt','.txt'}
GREEN='#173D2A';MID='#2E6A49';MINT='#DDF3E6';INK='#17211B';MUTED='#647269';GOLD='#C99A46'

def safe_target(user,job,filename,format_name):
    format_name=str(format_name or 'pdf').casefold().lstrip('.')
    if format_name not in {'pdf','docx','md','txt'}:raise Denied('document_format','Desteklenen biçimler: PDF, DOCX, Markdown ve TXT.',422)
    suffix='.'+format_name;name=Path(str(filename or 'belge'+suffix)).name
    if not name.casefold().endswith(suffix):name+=suffix
    name=re.sub(r'[^\w.()\- ]','-',name,flags=re.UNICODE).strip(' .')[:160]
    if not name:raise Denied('document_filename','Geçerli bir dosya adı gerekli.',422)
    base=ROOT/'users'/str(user['id'])/'personal-worker'/str(user.get('workspace_epoch',0))/'workspace';base.mkdir(parents=True,exist_ok=True)
    return base/name,format_name

def lines(content):return [line.rstrip() for line in str(content or '').replace('\r','').split('\n')]

def without_source_appendix(content,sources):
    rows=lines(content)
    if not sources:return rows
    for index,raw in enumerate(rows):
        if index>=len(rows)//2 and re.match(r'^\s*(?:#{1,6}\s*)?(?:kaynakça|kaynaklar|references)\b',raw,re.I):return rows[:index]
    return rows

def is_long_report(content,sources):
    words=len(re.findall(r'\b\w+\b',str(content),re.UNICODE))
    entries=len(re.findall(r'(?m)^\s*(?:#{1,4}\s*)?\d{1,3}[.)\]:-]\s+\S',str(content)))
    return words>=700 or entries>=8 or len(sources)>=8

def plain_summary(content,limit=430):
    parts=[]
    for raw in lines(content):
        value=raw.strip()
        if not value or re.match(r'^(?:#{1,6}\s+|[-*•]\s+|\d+[.)]\s+|https?://)',value):continue
        parts.append(value)
        if len(' '.join(parts))>=limit:break
    text=re.sub(r'\*\*|__|`','',' '.join(parts));text=re.sub(r'\s+',' ',text).strip()
    if len(text)<=limit:return text
    clipped=text[:limit];boundary=max(clipped.rfind('. '),clipped.rfind('! '),clipped.rfind('? '))
    return clipped[:boundary+1] if boundary>limit*.42 else clipped.rsplit(' ',1)[0]+'…'

def rich(text):
    """Small safe Markdown subset for ReportLab paragraphs."""
    escaped=html.escape(str(text),quote=True)
    escaped=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',escaped)
    escaped=re.sub(r'`([^`]+)`',r'<font name="InovensMono">\1</font>',escaped)
    def link(match):
        raw=html.unescape(match.group(0)).rstrip('.,;)');suffix=match.group(0)[len(html.escape(raw,quote=True)):]
        host=urllib.parse.urlsplit(raw).netloc.removeprefix('www.') or 'kaynak'
        return f'<link href="{html.escape(raw,quote=True)}" color="{MID}"><u>{html.escape(host)}</u></link>'+suffix
    return re.sub(r'https?://[^\s&lt;&gt;]+',link,escaped)

def pdf(path,title,content,sources):
    from reportlab.lib.enums import TA_CENTER,TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,HRFlowable,KeepTogether,Table,TableStyle
    regular='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf';bold='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf';mono='/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'
    for name,file in [('Inovens',regular),('InovensBold',bold),('InovensMono',mono)]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,file))
    body=ParagraphStyle('Body',fontName='Inovens',fontSize=9.6,leading=14.7,textColor=HexColor(INK),spaceAfter=6.5,allowWidows=0,allowOrphans=0)
    lead=ParagraphStyle('Lead',parent=body,fontSize=11.2,leading=17.2,textColor=HexColor('#304038'))
    title_style=ParagraphStyle('Title',parent=body,fontName='InovensBold',fontSize=27,leading=34,textColor=HexColor(GREEN),spaceAfter=12)
    eyebrow=ParagraphStyle('Eyebrow',parent=body,fontName='InovensBold',fontSize=9,leading=12,textColor=HexColor(MID),spaceAfter=9,tracking=1.2)
    h1=ParagraphStyle('H1',parent=body,fontName='InovensBold',fontSize=17,leading=22,textColor=HexColor(GREEN),spaceBefore=15,spaceAfter=8,keepWithNext=True)
    h2=ParagraphStyle('H2',parent=body,fontName='InovensBold',fontSize=12.5,leading=17,textColor=HexColor(MID),spaceBefore=11,spaceAfter=6,keepWithNext=True)
    h3=ParagraphStyle('H3',parent=body,fontName='InovensBold',fontSize=10.3,leading=14.5,textColor=HexColor(INK),spaceBefore=8,spaceAfter=4,keepWithNext=True)
    item=ParagraphStyle('Item',parent=body,leftIndent=12,firstLineIndent=-8,spaceAfter=5)
    source_style=ParagraphStyle('Source',parent=body,fontSize=8.4,leading=12.5,textColor=HexColor(MUTED),leftIndent=12,firstLineIndent=-8,spaceAfter=4)
    cover_small=ParagraphStyle('CoverSmall',parent=body,fontSize=8.5,leading=12,textColor=HexColor(MUTED))
    long=is_long_report(content,sources);story=[]
    if long:
        story += [Spacer(1,24*mm),Paragraph("INOVENS'AI · ARAŞTIRMA",eyebrow),Paragraph(rich(title or path.stem),title_style),
                  HRFlowable(width='34%',thickness=2,color=HexColor(GOLD),hAlign='LEFT',spaceBefore=3,spaceAfter=15),
                  Paragraph(rich(plain_summary(content)),lead),Spacer(1,16*mm)]
        stats=[]
        entry_count=len(re.findall(r'(?m)^\s*(?:#{1,4}\s*)?\d{1,3}[.)\]:-]\s+\S',str(content)))
        if entry_count:stats.append([Paragraph(f'<b>{entry_count}</b><br/><font size="7">İNCELENEN KAYIT</font>',cover_small)])
        stats.append([Paragraph(f'<b>{len(sources)}</b><br/><font size="7">KAYNAK</font>',cover_small)])
        stats.append([Paragraph(f'<b>{date.today().strftime("%d.%m.%Y")}</b><br/><font size="7">RAPOR TARİHİ</font>',cover_small)])
        table=Table([sum(stats,[])],colWidths=[48*mm]*len(stats),hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor('#F2F7F4')),('BOX',(0,0),(-1,-1),.5,HexColor('#CADDD1')),('INNERGRID',(0,0),(-1,-1),.5,HexColor('#CADDD1')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
        story += [table,Spacer(1,24*mm),Paragraph("Hazırlayan  INOVENS'AI",cover_small),PageBreak()]
    else:story += [Paragraph("INOVENS'AI",eyebrow),Paragraph(rich(title or path.stem),title_style),HRFlowable(width='100%',thickness=.7,color=HexColor('#CADDD1'),spaceAfter=12)]
    numbered=0
    for raw in without_source_appendix(content,sources):
        value=raw.strip()
        if not value:story.append(Spacer(1,3));continue
        if value=='---':story.append(HRFlowable(width='100%',thickness=.5,color=HexColor('#D8E3DC'),spaceBefore=6,spaceAfter=8));continue
        heading=re.match(r'^(#{1,6})\s+(.+)$',value)
        if heading:
            level=len(heading.group(1));label=heading.group(2)
            if long and re.search(r'^(?:kaynaklar|references)\b',label,re.I) and story and not isinstance(story[-1],PageBreak):story.append(PageBreak())
            story.append(Paragraph(rich(label),h1 if level==1 else h2 if level==2 else h3));continue
        numbered_match=re.match(r'^(\d{1,3})[.)\]:-]\s+(.+)$',value)
        if numbered_match and long and len(value)<180:
            numbered+=1;story.append(KeepTogether([Paragraph(f'<font color="{GOLD}">{int(numbered_match.group(1)):02d}</font>  '+rich(numbered_match.group(2)),h2),HRFlowable(width='100%',thickness=.35,color=HexColor('#D8E3DC'),spaceAfter=5)]));continue
        if re.match(r'^[-*•]\s+',value):story.append(Paragraph('• '+rich(re.sub(r'^[-*•]\s+','',value)),item));continue
        if numbered_match:story.append(Paragraph(rich(value),item));continue
        story.append(Paragraph(rich(value),body))
    if sources:
        if long:story.append(PageBreak())
        story += [Paragraph('Kaynakça',h1),Paragraph('Rapor hazırlanırken açılan ve içerikte kullanılan bağlantılar.',body)]
        for index,source in enumerate(sources[:80],1):story.append(Paragraph(f'{index:02d} · '+rich(str(source)),source_style))
    def page(canvas,doc):
        canvas.saveState();width,height=A4
        if doc.page>1 or not long:
            canvas.setStrokeColor(HexColor('#D8E3DC'));canvas.setLineWidth(.5);canvas.line(18*mm,height-14*mm,width-18*mm,height-14*mm)
            canvas.setFont('InovensBold',7.5);canvas.setFillColor(HexColor(MID));canvas.drawString(18*mm,height-10.5*mm,"INOVENS'AI")
            canvas.setFont('Inovens',7.2);canvas.setFillColor(HexColor(MUTED));canvas.drawRightString(width-18*mm,height-10.5*mm,str(title or path.stem)[:70])
        canvas.setFont('Inovens',7.4);canvas.setFillColor(HexColor(MUTED));canvas.drawString(18*mm,10*mm,'Doğrulanmış belge çıktısı');canvas.drawRightString(width-18*mm,10*mm,f'{doc.page:02d}')
        canvas.restoreState()
    document=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=19*mm,bottomMargin=17*mm,
                               title=str(title or path.stem),author="INOVENS'AI",subject='Araştırılmış ve kalite kontrolünden geçmiş belge',keywords='INOVENS AI, araştırma, rapor')
    document.build(story,onFirstPage=page,onLaterPages=page)

def docx(path,title,content,sources):
    from docx import Document
    from docx.shared import Pt,RGBColor,Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document=Document();section=document.sections[0];section.top_margin=Inches(.75);section.bottom_margin=Inches(.7);section.left_margin=Inches(.8);section.right_margin=Inches(.8)
    normal=document.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(23,33,27)
    for style_name,size,color in [('Title',26,GREEN),('Heading 1',17,GREEN),('Heading 2',13,MID),('Heading 3',11,INK)]:
        style=document.styles[style_name];style.font.name='DejaVu Sans';style.font.size=Pt(size);style.font.bold=True;style.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
    header=section.header.paragraphs[0];header.text="INOVENS'AI";header.style=document.styles['Caption']
    footer=section.footer.paragraphs[0];footer.text='Doğrulanmış belge çıktısı';footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.core_properties.title=str(title or path.stem);document.core_properties.author="INOVENS'AI";document.core_properties.subject='Kalite kontrolünden geçmiş belge'
    document.add_heading(str(title or path.stem),0)
    if is_long_report(content,sources):document.add_paragraph(plain_summary(content),style='Subtitle');document.add_page_break()
    for raw in lines(content):
        value=raw.strip()
        if not value:continue
        heading=re.match(r'^(#{1,6})\s+(.+)$',value)
        if heading:document.add_heading(heading.group(2),level=min(3,len(heading.group(1))));continue
        if re.match(r'^[-*•]\s+',value):document.add_paragraph(re.sub(r'^[-*•]\s+','',value),style='List Bullet');continue
        if re.match(r'^\d+[.)]\s+',value):document.add_paragraph(value,style='List Number');continue
        document.add_paragraph(value)
    if sources:
        document.add_page_break();document.add_heading('Kaynakça',level=1)
        for source in sources[:80]:document.add_paragraph(str(source),style='List Bullet')
    document.save(path)

def create(user,job,args):
    title=str(args.get('title','')).strip()[:240];content=str(args.get('content','')).strip();sources=args.get('sources',[])
    if len(content)<20 or len(content)>500_000:raise Denied('document_content','Belge içeriği boş, çok kısa veya fazla uzun.',422)
    if not isinstance(sources,list):raise Denied('document_sources','Kaynaklar liste olmalıdır.',422)
    path,format_name=safe_target(user,job,args.get('filename'),args.get('format'))
    if format_name=='pdf':pdf(path,title,content,sources)
    elif format_name=='docx':docx(path,title,content,sources)
    else:
        rendered=(('# '+title+'\n\n') if title and format_name=='md' else (title+'\n\n' if title else ''))+content
        if sources:rendered+='\n\nKaynaklar\n'+'\n'.join('- '+str(x) for x in sources[:80])
        path.write_text(rendered,encoding='utf-8')
    if not path.is_file() or path.stat().st_size<20:raise Denied('document_create_failed','Belge oluşturulamadı.',500)
    qa={'ok':True,'screenshots':[]}
    if format_name in {'pdf','docx'}:
        from document_qa import validate_document
        qa=validate_document(path,path.parent/'.qa')
        if not qa['ok']:
            path.unlink(missing_ok=True);raise Denied('document_quality_failed','Belge okunabilirlik kontrolünden geçmedi; içeriği veya yerleşimi düzeltip yeniden deneyin.',422)
    return {'path':'/workspace/'+path.name,'filename':path.name,'format':format_name,'size':path.stat().st_size,
            'pages':qa.get('pages'),'quality_check':'passed','verified':True}
